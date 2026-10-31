# -*- coding: utf-8 -*-
"""自动化回归测试（自包含，不依赖本机任何文件）。

运行：python test_core.py
全部通过时最后打印 ALL PASS；有失败会列出具体项目并返回退出码 1。

样本数据全部在临时目录里现场生成：
- 一张仿真实结构的德育分统计表（第3行大类、第4行细分、数据从第5行）
- 一个 5 列 Word 模板
"""
import os
import sys
import shutil
import tempfile
import traceback

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from docx import Document

import main

TMP = None
RESULTS = []


# ---------------------------------------------------------------- 测试框架
def check(name, func):
    """执行一个检查项，记录通过/失败。"""
    try:
        func()
        RESULTS.append((name, True, ""))
        print(f"  [OK]   {name}")
    except Exception as exc:
        RESULTS.append((name, False, str(exc)))
        print(f"  [FAIL] {name}")
        print(f"         {exc}")
        if os.environ.get("TEST_VERBOSE"):
            traceback.print_exc()


def eq(actual, expected, what=""):
    if actual != expected:
        raise AssertionError(f"{what} 期望 {expected!r}，实际 {actual!r}")


def truth(value, what=""):
    if not value:
        raise AssertionError(f"{what} 期望为真，实际 {value!r}")


# ---------------------------------------------------------------- 样本构造
def build_sample_xlsx(path):
    """构造一张与真实统计表结构一致的样本：
    A=姓名 B=基础分 C=基础分鉴定 D=3-4月加分描述 E=3-4月表扬信 F=鉴定
    G/H=扣分 I=5-8月加分描述 J=5-8月表扬信 K=鉴定 L/M=扣分 N=备注 O=序号 P=德育总分
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "大三下"

    ws["A1"] = "2025-2026学年第2学期德育分统计表"
    ws.merge_cells("A1:P1")
    ws["A2"] = "班级：农学23-1班"
    ws.merge_cells("A2:B2")

    # 第 3 行：大类（合并单元格，和真实表一致）
    for rng, text in (("A3:A4", "姓名"), ("B3:C3", "基础分"),
                      ("D3:F3", "3-4月加分项"), ("G3:H3", "3-4月扣分项"),
                      ("I3:K3", "5-8月加分项"), ("L3:M3", "5-8月扣分项"),
                      ("N3:N4", "备注"), ("O3:O4", "序号"), ("P3:P4", "德育总分")):
        ws[rng.split(":")[0]] = text
        ws.merge_cells(rng)

    # 第 4 行：细分
    for cell, text in (("C4", "鉴定"), ("D4", "3-4月"), ("E4", "表扬信"),
                       ("F4", "鉴定"), ("G4", "鉴定"), ("H4", "鉴定"),
                       ("I4", "5-8月"), ("J4", "表扬信"), ("K4", "鉴定"),
                       ("L4", "鉴定"), ("M4", "鉴定")):
        ws[cell] = text

    # 第 5 行起：数据
    ws.append(["张三", "基础分+72", "72", "院3月学霸宿舍奖+2",
               "3.6志愿服务表扬信+0.5\n3.10表彰大会表扬信+0.5", "", "", "", "",
               "7.5宿舍搬迁表扬信+0.5", "", "", "", "班长", "1", 82])
    ws.append(["李四", "基础分+75", "75", "", "", "", "", "", "", "", "", "",
               "", "", "2", None])
    ws.append(["王五", "基础分+60", "60", "", "3.8器材搬运表扬信+0.5", "", "", "", "",
               "6.5反诈活动表扬信+0.5", "", "", "", "", "3", -5])
    # 干扰行：不该被当成学生
    ws.append(["合计", "", "207", "", "", "", "", "", "", "", "", "", "", "", "", 207])
    wb.save(path)


def build_sample_template(path):
    """构造一个 5 列 Word 模板：标题行 + 表头行 + 空数据行。"""
    doc = Document()
    table = doc.add_table(rows=3, cols=5)
    cell = table.rows[0].cells[0]
    cell.text = "XX班 2025-2026 第二学期德育分确认签字表"
    cell.add_paragraph("班级")
    for i, head in enumerate(["姓名", "德育分", "表扬信", "确认签字", "备注"]):
        table.rows[1].cells[i].text = head
    doc.save(path)


# ---------------------------------------------------------------- 各项检查
def test_detect_columns():
    wb = openpyxl.load_workbook(os.path.join(TMP, "sample.xlsx"), data_only=True)
    ws = wb.active
    headers = main.scan_headers(ws)
    name_col, score_col, praise_cols = main.detect_columns(headers)
    wb.close()
    eq(name_col, 1, "姓名列")
    eq(score_col, 16, "总分列")
    eq(praise_cols, [5, 10], "表扬信列")


def test_month_groups():
    wb = openpyxl.load_workbook(os.path.join(TMP, "sample.xlsx"), data_only=True)
    ws = wb.active
    headers = main.scan_headers(ws)
    _, _, praise_cols = main.detect_columns(headers)
    groups = main.group_praise_by_month(ws, praise_cols)
    wb.close()
    eq(groups, [("3-4月", [5]), ("5-8月", [10])], "月份分组")


def test_title_and_class():
    wb = openpyxl.load_workbook(os.path.join(TMP, "sample.xlsx"), data_only=True)
    ws = wb.active
    title = main.extract_title(ws, "大三下")
    cls = main.extract_class(ws)
    wb.close()
    eq(title, "大三下 2025-2026 第二学期德育分确认签字表", "标题")
    eq(cls, "农学23-1班", "班级")


def test_extract_students():
    wb = openpyxl.load_workbook(os.path.join(TMP, "sample.xlsx"), data_only=True)
    ws = wb.active
    students = main.extract_students(ws, 1, 16, [5, 10])
    wb.close()
    by = {s["name"]: s for s in students}
    eq(len(students), 3, "学生人数（合计行应被排除）")
    eq(by["张三"]["praise"], 3, "张三表扬信条数")
    eq(by["李四"]["praise"], 0, "李四表扬信条数")
    eq(by["王五"]["praise"], 2, "王五表扬信条数")
    eq(by["张三"]["score"], "82", "张三分数字符串")


def test_validate_students():
    wb = openpyxl.load_workbook(os.path.join(TMP, "sample.xlsx"), data_only=True)
    ws = wb.active
    students = main.extract_students(ws, 1, 16, [5, 10])
    wb.close()
    problems = main.validate_students(students)
    kinds = {(p[0], p[2]) for p in problems}
    eq(len(problems), 2, "问题条数（李四空分数 + 王五负分）")
    truth(("empty", "李四") in kinds, "应检出李四分数为空")
    truth(("neg", "王五") in kinds, "应检出王五负分")


def test_cert_and_praise_extract():
    data = main.extract_persons_from_workbook(os.path.join(TMP, "sample.xlsx"))
    truth("张三" in data, "应提取到张三")
    certs, praises = data["张三"]
    eq(len(certs), 1, "张三证书条数")
    eq(certs[0], "院3月学霸宿舍奖", "张三证书内容（应去掉 +2）")
    eq(len(praises), 3, "张三表扬信条数（跨两个月份段合并）")


def test_add_year_to_date():
    sem = "2025-2026学年第2学期"
    eq(main.add_year_to_date("3.6志愿服务", sem), "2026年3月6日志愿服务", "补年份-月.日")
    eq(main.add_year_to_date("五月劳动周", sem), "2026年5月劳动周", "补年份-中文月份")
    eq(main.add_year_to_date("9月义务劳动", "2025-2026学年第一学期"),
       "2025年9月义务劳动", "第一学期 9 月用起始年")
    eq(main.add_year_to_date("3.6志愿服务", "某班统计表"),
       "3.6志愿服务", "无法解析学期时原样返回")


def test_cert_sentences():
    eq(main.to_cert_sentence("2026年3月参与发表论文"),
       "2026年3月参与发表论文", "参与类不加荣获")
    eq(main.to_cert_sentence("2026年6月省志愿者证书"),
       "2026年6月被评为省志愿者", "身份类用被评为")
    eq(main.to_cert_sentence("2026年3月校一等奖"),
       "2026年3月荣获校一等奖", "获奖类用荣获")
    eq(main.to_cert_sentence("2026年3月优秀团员"),
       "2026年3月被评为优秀团员", "荣誉称号用被评为")


def test_praise_sentences():
    eq(main.to_praise_sentence("2026年7月5日宿舍搬迁表扬信"),
       "2026年7月5日参与宿舍搬迁活动", "表扬信加参与与活动")
    eq(main.to_praise_sentence("2026年6月年级大会表扬信"),
       "2026年6月参与年级大会", "会议类不加活动")
    eq(main.to_praise_sentence("2026年6月参与义务劳动表扬信"),
       "2026年6月参与义务劳动活动", "不重复参与")


def test_parse_name_text():
    names, dup = main.parse_name_text("张三\n李四\n张三\n张三（男）\n李四")
    eq(names, ["张三", "李四"], "去重后的名单")
    eq(dup, 3, "被忽略的重复个数")
    names, dup = main.parse_name_text("甲，乙、丙;丁 戊")
    eq(names, ["甲", "乙", "丙", "丁", "戊"], "多种分隔符")
    eq(dup, 0, "无重复时计数为 0")


def test_collect_and_write_then_convert():
    xlsx = os.path.join(TMP, "农学23-1班2025-2026学年第2学期德育分统计表.xlsx")
    rows, missing, multi = main.collect_person_rows([xlsx], ["张三", "李四", "查无此人"])
    eq(len(rows), 2, "结果行数")
    eq(missing, ["查无此人"], "未找到的人员")
    eq(multi, [], "多源表人员")

    out1 = os.path.join(TMP, "提取结果.xlsx")
    main.write_person_result_xlsx(rows, out1)
    wb = openpyxl.load_workbook(out1)
    ws = wb.active
    eq([c.value for c in ws[1]], ["姓名", "证书", "表扬信", "来源学期"], "结果表头")
    eq(ws.cell(2, 4).value, "2025-2026学年第2学期", "来源学期列")
    wb.close()

    out2 = os.path.join(TMP, "提取结果_句式.xlsx")
    n = main.convert_result_to_sentences(out1, out2)
    eq(n, 2, "句式转换行数")
    wb = openpyxl.load_workbook(out2)
    ws = wb.active
    eq(ws.cell(2, 4).value, "2025-2026学年第2学期", "句式转换后仍保留来源学期")
    body = str(ws.cell(2, 2).value or "")
    truth("荣获" in body or "参与" in body, "证书应已转成句子")
    wb.close()


def test_template_structure():
    ok, msg = main.check_template_structure(os.path.join(TMP, "template.docx"))
    truth(ok, f"模板校验应通过（{msg}）")
    ok, msg = main.check_template_structure(os.path.join(TMP, "不存在.docx"))
    truth(not ok, "不存在的模板应校验失败")


def test_fill_word():
    tpl = os.path.join(TMP, "template.docx")
    out = os.path.join(TMP, "签字表.docx")
    wb = openpyxl.load_workbook(os.path.join(TMP, "sample.xlsx"), data_only=True)
    ws = wb.active
    students = main.extract_students(ws, 1, 16, [5, 10])
    title = main.extract_title(ws, "大三下")
    cls = main.extract_class(ws)
    wb.close()

    main.fill_word(tpl, out, students, title, cls)
    table = Document(out).tables[0]
    eq(len(table.rows), 5, "生成表格行数（标题+表头+3人）")
    paras = [p.text for p in table.rows[0].cells[0].paragraphs]
    eq(paras[0], title, "标题行")
    truth("班级：农学23-1班" in paras, "班级应写入标题行")
    eq([c.text for c in table.rows[2].cells], ["张三", "82", "3", "", ""], "第一行数据")
    eq([c.text for c in table.rows[4].cells], ["王五", "-5", "2", "", ""], "第三行数据")


def test_unique_path():
    base = os.path.join(TMP, "同名.xlsx")
    eq(main.unique_path_on_disk(base), base, "文件不存在时保持原名")
    open(base, "w").close()
    second = main.unique_path_on_disk(base)
    eq(os.path.basename(second), "同名_1.xlsx", "已存在时加 _1")
    open(second, "w").close()
    eq(os.path.basename(main.unique_path_on_disk(base)), "同名_2.xlsx", "再次冲突时加 _2")


def test_config_roundtrip():
    cfg_path = os.path.join(TMP, "config.json")
    orig = main.config_file
    main.config_file = lambda: cfg_path
    try:
        eq(main.load_config(), {}, "无配置时返回空字典")
        truth(main.save_config({"template_path": r"C:\t.docx", "output_dir": TMP}),
              "保存应成功")
        cfg = main.load_config()
        eq(cfg["template_path"], r"C:\t.docx", "模板路径往返")
        eq(cfg["output_dir"], TMP, "输出目录往返")
    finally:
        main.config_file = orig


def test_process_one_excel():
    outdir = os.path.join(TMP, "out")
    os.makedirs(outdir, exist_ok=True)
    xlsx = os.path.join(TMP, "农学23-1班2025-2026学年第2学期德育分统计表.xlsx")
    out, n = main.process_one_excel(xlsx, os.path.join(TMP, "template.docx"), outdir)
    eq(n, 3, "生成人数")
    truth(os.path.exists(out), "输出文件应存在")
    eq(os.path.basename(out),
       "农学23-1班2025-2026学年第2学期德育分签字表.docx", "输出文件名")


def test_process_one_excel_month_filter():
    """只统计 3-4月：张三应剩 2 条表扬信。"""
    outdir = os.path.join(TMP, "out2")
    os.makedirs(outdir, exist_ok=True)
    xlsx = os.path.join(TMP, "农学23-1班2025-2026学年第2学期德育分统计表.xlsx")
    out, n = main.process_one_excel(
        xlsx, os.path.join(TMP, "template.docx"), outdir, selected_months={"3-4月"})
    eq(n, 3, "人数")
    rows = Document(out).tables[0].rows
    eq(rows[2].cells[2].text, "2", "张三只统计 3-4月 后的表扬信条数")


# ---------------------------------------------------------------- 主流程
def main_entry():
    global TMP
    TMP = tempfile.mkdtemp(prefix="moral_test_")
    try:
        build_sample_xlsx(os.path.join(TMP, "sample.xlsx"))
        build_sample_template(os.path.join(TMP, "template.docx"))
        # 提取流程按文件名解析学期，另存一份带学期的文件名
        shutil.copy(os.path.join(TMP, "sample.xlsx"),
                    os.path.join(TMP, "农学23-1班2025-2026学年第2学期德育分统计表.xlsx"))

        print("\n=== 表头识别与数据提取 ===")
        check("识别姓名/总分/表扬信列", test_detect_columns)
        check("月份段分组", test_month_groups)
        check("标题与班级识别", test_title_and_class)
        check("学生数据提取（合计行不应混入）", test_extract_students)
        check("分数校验（空分数/负分）", test_validate_students)

        print("\n=== 证书 / 表扬信 ===")
        check("跨月份提取证书与表扬信", test_cert_and_praise_extract)
        check("日期补年份", test_add_year_to_date)
        check("证书句子措辞", test_cert_sentences)
        check("表扬信句子措辞", test_praise_sentences)

        print("\n=== 名单与结果文件 ===")
        check("名单解析与去重", test_parse_name_text)
        check("提取结果写出 / 句式转换保留来源学期", test_collect_and_write_then_convert)
        check("同名文件自动另存 _1/_2", test_unique_path)

        print("\n=== Word 生成 ===")
        check("模板结构校验", test_template_structure)
        check("灌入学生数据", test_fill_word)
        check("端到端生成（含文件名替换）", test_process_one_excel)
        check("按月份段生成", test_process_one_excel_month_filter)

        print("\n=== 配置 ===")
        check("配置读写往返", test_config_roundtrip)

        passed = sum(1 for _, ok, _ in RESULTS if ok)
        total = len(RESULTS)
        print(f"\n{passed}/{total} 项通过")
        if passed == total:
            print("ALL PASS")
            return 0
        print("\n失败项：")
        for name, ok, msg in RESULTS:
            if not ok:
                print(f"  - {name}: {msg}")
        return 1
    finally:
        shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main_entry())
