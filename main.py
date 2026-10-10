# -*- coding: utf-8 -*-
"""
德育分签字表生成器
- 不写死文件名 / 工作表 / 表头名 / 列位置 / 人数，通用适配。
- 表头自动识别（姓名 / 总分 / 表扬信），可手动调整。
- 表扬信按“月份段”分组，可只选某几个月份段；按条目条数累加，没有填 0。
- Word 标题自动识别为“年级 学年 第X学期德育分确认签字表”。
- 输出文件名：数据源名中的“统计表”替换为“签字表”。
- 模板路径、固定输出目录通过 config.json 记住，下次打开不重置。
"""
import os
import re
import sys
import json
import copy
import threading
import webbrowser
import urllib.request

import openpyxl
from openpyxl.utils import get_column_letter

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from tkinterdnd2 import TkinterDnD, DND_FILES

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

HEADER_SCAN_ROWS = 8
NAME_KW = ["姓名", "名字", "学生姓名", "name"]
SCORE_KW = ["总分", "合计", "总成绩", "总得分", "最终得分", "德育总分", "德育分"]
PRAISE_KW = ["表扬信", "通报表扬", "表扬"]
ROW_BLACK = ["合计", "总计", "小计", "说明", "备注", "签字", "序号", "编号",
             "日期", "制表", "审核", "姓名", "班级", "统计表", "评分", "汇总",
             "一览", "学年", "学期", "表格"]
TEMPLATE_NAME = "签字表模板.docx"
CONFIG_NAME = "config.json"
APP_DIR_NAME = "德育分签字表生成器"
APP_VERSION = "v1.6.0"
GITHUB_URL = "https://github.com/yyxi746/moral-score-sign-generator"
CN_NUM = {"1": "一", "2": "二", "3": "三", "4": "四",
          "一": "一", "二": "二", "三": "三", "四": "四"}


# ============================ 通用辅助 ============================
def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def resource_path(name):
    """打包后资源在临时解压目录，开发时在脚本目录。"""
    if getattr(sys, "frozen", False):
        return os.path.join(getattr(sys, "_MEIPASS", app_dir()), name)
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), name)


def is_number_text(text):
    try:
        float(str(text).strip())
        return True
    except (ValueError, TypeError):
        return False


def fmt_score(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        f = float(v)
        return str(int(f)) if f.is_integer() else str(f)
    s = str(v).strip()
    if is_number_text(s):
        f = float(s)
        return str(int(f)) if f.is_integer() else str(f)
    return s


def match_any(header, keywords):
    h = str(header).lower()
    return any(k.lower() in h for k in keywords)


# ============================ 配置持久化 ============================
def app_config_dir():
    """系统隐藏目录（%APPDATA%\\德育分签字表生成器），EXE 放哪都不影响。"""
    base = os.environ.get("APPDATA") or os.path.expanduser("~")
    return os.path.join(base, APP_DIR_NAME)


def config_file():
    d = app_config_dir()
    try:
        os.makedirs(d, exist_ok=True)
    except Exception:
        pass
    return os.path.join(d, CONFIG_NAME)


def load_config():
    try:
        with open(config_file(), "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_config(cfg):
    try:
        with open(config_file(), "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


# ============================ Excel 读取与识别 ============================
def list_sheets(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def scan_headers(ws):
    max_col = ws.max_column or 1
    max_row = ws.max_row or 1
    scan = min(HEADER_SCAN_ROWS, max_row)
    headers = {}
    for c in range(1, max_col + 1):
        parts = []
        for r in range(1, scan + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                parts.append(str(v).strip())
        headers[c] = " ".join(parts)
    return headers


def column_labels(headers):
    out = []
    for c in sorted(headers):
        letter = get_column_letter(c)
        h = headers[c].replace("\n", " ")
        if len(h) > 18:
            h = h[:18] + "…"
        out.append((f"{letter}｜{h}", c))
    return out


def detect_columns(headers):
    cols = sorted(headers)
    name_cands = [c for c in cols if match_any(headers[c], NAME_KW)]
    score_cands = [c for c in cols
                   if match_any(headers[c], SCORE_KW) and "基础分" not in headers[c]]
    praise_cands = [c for c in cols if match_any(headers[c], PRAISE_KW)]
    name_col = name_cands[0] if name_cands else None
    score_col = score_cands[-1] if score_cands else None
    return name_col, score_col, praise_cands


def extract_class(ws):
    max_row = ws.max_row or 1
    scan = min(HEADER_SCAN_ROWS, max_row)
    for r in range(1, scan + 1):
        for c in range(1, (ws.max_column or 1) + 1):
            v = ws.cell(r, c).value
            if v and "班级" in str(v):
                t = str(v)
                for sep in ["：", ":"]:
                    if sep in t:
                        name = t.split(sep, 1)[1].strip()
                        if name:
                            return name
    return ""


# ---------- 表扬信：按月份段分组 ----------
def find_month_group(ws, col):
    total_cols = ws.max_column or 1
    cands = []
    for rng in ws.merged_cells.ranges:
        if rng.min_col <= col <= rng.max_col and rng.min_row <= HEADER_SCAN_ROWS:
            v = ws.cell(rng.min_row, rng.min_col).value
            if v:
                t = str(v).strip()
                # 排除横跨整表的大标题
                if (rng.max_col - rng.min_col + 1) >= total_cols - 1:
                    continue
                cands.append(t)
    # 优先：合并标题中含具体数字月份
    for t in cands:
        if re.search(r"\d+\s*月", t):
            return t
    # 其次：表头行内从该列向左找含“月”的标题（覆盖合并锚点/非合并子标题）
    for r in range(1, HEADER_SCAN_ROWS + 1):
        for cc in range(col, 0, -1):
            v = ws.cell(r, cc).value
            if not v:
                continue
            t = str(v).strip()
            if "月" in t and "统计表" not in t and "鉴定" not in t and t != "表扬信":
                return t
    # 合并候选里含“月”但无数字（如 X-X月加分项）
    for t in cands:
        if "月" in t:
            return t
    return cands[0] if cands else None


def month_of_group(text):
    m = re.search(r"(\d+\s*[-—~]\s*\d+\s*月)", text)
    if m:
        return m.group(1).replace(" ", "")
    m = re.search(r"([XxX]\s*[-—~]\s*[XxX]\s*月)", text)
    if m:
        return m.group(1).replace(" ", "")
    m = re.search(r"(\d+\s*月)", text)
    return m.group(1) if m else text


def group_praise_by_month(ws, praise_cols):
    groups, order = {}, []
    for c in praise_cols:
        g = find_month_group(ws, c)
        month = month_of_group(g) if g else (get_column_letter(c) + "列")
        if month not in groups:
            groups[month] = []
            order.append(month)
        groups[month].append(c)
    return [(m, groups[m]) for m in order]


# 月份段跨学年的标准先后顺序
MONTH_ORDER = ["9-10月", "10-11月", "11-12月", "12-1月", "1-2月",
               "2-3月", "3-4月", "4-5月", "5-6月", "5-8月",
               "6-7月", "7-8月"]


def month_sort_key(m):
    for i, k in enumerate(MONTH_ORDER):
        if k in m:
            return i
    mm = re.search(r"(\d+)", m)
    return 100 + (int(mm.group(1)) if mm else 0)


# ---------- 标题信息 ----------
def extract_title(ws, sheet_name):
    parts_text = []
    max_row = ws.max_row or 1
    for r in range(1, min(HEADER_SCAN_ROWS, max_row) + 1):
        for c in range(1, (ws.max_column or 1) + 1):
            v = ws.cell(r, c).value
            if v:
                parts_text.append(str(v))
    blob = " ".join(parts_text)
    ym = re.search(r"(20\d{2}\s*[-—~]\s*20\d{2})", blob)
    tm = re.search(r"第\s*([1-4一二三四])\s*学期", blob)
    term = ("第" + CN_NUM[tm.group(1)] + "学期") if tm else ""
    sheet_part = sheet_name or ""
    if "XX" in sheet_part:
        sheet_part = ""
    head = " ".join(p for p in [sheet_part,
                                ym.group(1) if ym else "", term] if p)
    return head + "德育分确认签字表"


# ---------- 数据行 ----------
def is_valid_name(text):
    if not text:
        return False
    text = str(text).strip()
    if not text or is_number_text(text):
        return False
    if any(k in text for k in ROW_BLACK):
        return False
    return True


def count_praise_cell(v):
    if v is None:
        return 0
    return sum(1 for line in str(v).split("\n") if line.strip())


def extract_students(ws, name_col, score_col, praise_cols):
    students = []
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    if not name_col or name_col > max_col:
        return students
    for r in range(1, max_row + 1):
        nv = ws.cell(r, name_col).value
        name = "" if nv is None else str(nv).strip()
        if not is_valid_name(name):
            continue
        score = fmt_score(ws.cell(r, score_col).value) if score_col and score_col <= max_col else ""
        praise = 0
        for pc in praise_cols:
            if pc and pc <= max_col:
                praise += count_praise_cell(ws.cell(r, pc).value)
        students.append({"name": name, "score": score, "praise": praise})
    return students


def validate_students(students):
    """生成前校验。返回问题列表：
    ("zero",) 没有学生；
    ("empty", 行号, 姓名) 分数为空；
    ("neg", 行号, 姓名, 分数) 分数为负；
    ("badtext", 行号, 姓名, 分数) 分数不是数字。"""
    problems = []
    if not students:
        return [("zero",)]
    for i, s in enumerate(students):
        score = s.get("score", "")
        if score == "" or score is None:
            problems.append(("empty", i, s.get("name", "")))
        else:
            try:
                if float(str(score).strip()) < 0:
                    problems.append(("neg", i, s.get("name", ""), score))
            except (ValueError, TypeError):
                problems.append(("badtext", i, s.get("name", ""), score))
    return problems


# ---------- 输出文件名 ----------
def output_file_name(excel_path):
    stem = os.path.splitext(os.path.basename(excel_path))[0]
    if "统计表" in stem:
        name = stem.replace("统计表", "签字表")
    else:
        name = stem + "签字表"
    return name + ".docx"


def compute_output_path(excel_path, output_dir):
    name = output_file_name(excel_path)
    d = output_dir.strip() if output_dir and output_dir.strip() else os.path.dirname(excel_path)
    return os.path.join(d, name)


def unique_output_path(path, reserved):
    """同一轮批量中若 path 已被别的源文件占用（同名不同源），自动加 _2/_3。"""
    if path not in reserved:
        return path
    base, ext = os.path.splitext(path)
    i = 2
    while f"{base}_{i}{ext}" in reserved:
        i += 1
    return f"{base}_{i}{ext}"


def ensure_writable(path):
    """输出文件已存在且被占用（如正在 Word 中打开）时给出友好提示。"""
    if os.path.exists(path):
        try:
            with open(path, "ab"):
                pass
        except PermissionError:
            raise PermissionError(
                "文件正在被使用（可能正在 Word 中打开），请关闭后重试："
                + os.path.basename(path))


# ============================ Word 填充 ============================
def set_run_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def fill_cell(cell, text, size=12):
    p = cell.paragraphs[0]
    run = p.add_run("" if text is None else str(text))
    set_run_font(run, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def replace_paragraph(p, text, size=18, bold=True):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def set_class_on_paragraph(p, class_name):
    t = p.text.strip()
    size = 12
    if p.runs and p.runs[0].font.size:
        size = p.runs[0].font.size.pt
    if t == "班级":
        run = p.add_run("：" + class_name)
        set_run_font(run, size=size)
    elif t in ("班级：", "班级:"):
        run = p.add_run(class_name)
        set_run_font(run, size=size)


def set_title_row(table, title, class_name):
    cell = table.rows[0].cells[0]
    for p in cell.paragraphs:
        t = p.text.strip()
        if t.startswith("班级"):
            if class_name:
                set_class_on_paragraph(p, class_name)
        elif "德育分确认签字表" in t or "签字表" in t:
            if title:
                replace_paragraph(p, title, size=18, bold=True)


def set_repeat_header(table):
    """让表头行（第 2 行）跨页时在每页顶部自动重复。"""
    try:
        if len(table.rows) >= 2:
            trPr = table.rows[1]._tr.get_or_add_trPr()
            if trPr.find(qn("w:tblHeader")) is None:
                el = OxmlElement("w:tblHeader")
                el.set(qn("w:val"), "true")
                trPr.append(el)
    except Exception:
        pass


def fill_word(template_path, output_path_, students, title, class_name):
    doc = Document(template_path)
    if not doc.tables:
        raise ValueError("模板中没有找到表格，请检查 Word 模板。")
    table = doc.tables[0]
    tbl = table._tbl
    all_tr = tbl.findall(qn("w:tr"))
    if len(all_tr) < 3:
        raise ValueError("模板表格结构异常（至少需要标题行、表头行、一条数据行）。")

    sample_tr = copy.deepcopy(all_tr[2])
    for tr in all_tr[2:]:
        tbl.remove(tr)

    for st in students:
        new_tr = copy.deepcopy(sample_tr)
        tbl.append(new_tr)
        row = table.rows[-1]
        values = [st["name"], st["score"], st["praise"], "", ""]
        for ci, val in enumerate(values):
            fill_cell(row.cells[ci], val, size=12)

    set_title_row(table, title, class_name)
    set_repeat_header(table)
    ensure_writable(output_path_)
    doc.save(output_path_)


# ============================ 批量处理 ============================
def collect_excels(folder):
    """递归收集文件夹下所有 xlsx/xlsm（排除 Excel 临时锁文件 ~$ 开头）。"""
    out = []
    for dirpath, _, files in os.walk(folder):
        for f in files:
            if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$"):
                out.append(os.path.join(dirpath, f))
    return sorted(out)


def analyze_excel_auto(path):
    """对单个文件自动识别（第一个工作表、全部月份组）。
    返回 dict：sheet/title/class/students；识别失败抛异常。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        if not wb.sheetnames:
            raise ValueError("工作簿中没有工作表。")
        sheet = wb.sheetnames[0]
        ws = wb[sheet]
        headers = scan_headers(ws)
        name_col, score_col, praise_cols = detect_columns(headers)
        if not name_col:
            raise ValueError("未识别到姓名列。")
        if not score_col:
            raise ValueError("未识别到德育分(总分)列。")
        month_groups = group_praise_by_month(ws, praise_cols)
        students = extract_students(ws, name_col, score_col, praise_cols)
        if not students:
            raise ValueError("未识别到学生数据。")
        return {"sheet": sheet,
                "title": extract_title(ws, sheet),
                "class": extract_class(ws),
                "students": students,
                "month_groups": month_groups}
    finally:
        wb.close()


def process_one_excel(path, template_path_, output_dir, selected_months=None,
                      overrides=None, out_path=None):
    """自动分析并生成一个 Word，返回 (输出路径, 人数)。
    selected_months：只统计这些月份段（集合）；None = 全部月份段。
    overrides：手工修正的标题/班级。
    out_path：显式指定输出路径（批量防同名覆盖时使用）。"""
    info = analyze_excel_auto(path)
    out = out_path or compute_output_path(path, output_dir)
    d = os.path.dirname(out)
    if d:
        os.makedirs(d, exist_ok=True)
    students = info["students"]
    if selected_months is not None:
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            ws = wb[info["sheet"]]
            headers = scan_headers(ws)
            name_col, score_col, praise_cols = detect_columns(headers)
            groups = group_praise_by_month(ws, praise_cols)
            sel_cols = []
            for month, cols in groups:
                if month in selected_months:
                    sel_cols.extend(cols)
            students = extract_students(ws, name_col, score_col, sel_cols)
        finally:
            wb.close()
    title = info["title"]
    cls = info["class"]
    if overrides:
        title = overrides.get("title", title)
        cls = overrides.get("class", cls)
    fill_word(template_path_, out, students, title, cls)
    return out, len(students)


# ============================ 图形界面 ============================
def register_drop_recursive(widget, handler):
    """递归为容器内所有控件注册文件拖放（不支持的控件自动跳过）。"""
    try:
        widget.drop_target_register(DND_FILES)
        widget.dnd_bind("<<Drop>>", handler)
    except Exception:
        pass
    for child in widget.winfo_children():
        register_drop_recursive(child, handler)


class App(TkinterDnD.Tk):
    def __init__(self):
        super().__init__()
        self.title("德育分签字表生成器")
        self.geometry("820x900")
        self.minsize(780, 700)
        ico = resource_path("app.ico")
        if os.path.exists(ico):
            try:
                self.iconbitmap(default=ico)
            except Exception:
                pass

        cfg = load_config()
        self._skip_version = cfg.get("skip_version", "")

        self.excel_path = tk.StringVar()
        tpl = cfg.get("template_path", "")
        init_tpl = tpl if (tpl and os.path.exists(tpl)) else os.path.join(app_dir(), TEMPLATE_NAME)
        self.template_path = tk.StringVar(value=init_tpl)
        self.output_dir = tk.StringVar(value=cfg.get("output_dir", ""))
        self.output_path = tk.StringVar()
        self.class_name = tk.StringVar()
        self.sheet_name = tk.StringVar()
        self.title_text = tk.StringVar()

        self.wb = None
        self.headers = {}
        self.col_options = []
        self.label_to_col = {}
        self.name_label = tk.StringVar()
        self.score_label = tk.StringVar()
        self.month_groups = []
        self.month_vars = {}

        # 批量处理状态
        self.batch_items = {}        # tree item -> 文件路径
        self.batch_info = {}         # tree item -> 分析结果 dict
        self.batch_path_set = set()  # 已加入文件去重
        self._analyze_queue = []
        self.batch_month_vars = {}   # 月份段 -> IntVar（批量页复选框）
        self.batch_overrides = {}    # item -> {"title":..,"class":..}
        self._gen_queue = []
        self._gen_results = {}
        self._gen_cancel = False
        self._gen_path_map = {}

        self._build_ui()

        if not os.path.exists(init_tpl):
            self.after(300, self._first_run_guide)
        self.after(1500, self._check_update)

    # ---------- 首次运行 / 模板引导 ----------
    def _first_run_guide(self):
        messagebox.showinfo(
            "欢迎使用",
            "欢迎使用德育分签字表生成器。\n\n"
            "首次使用需要选择一个 Word 签字表模板（.docx），\n"
            "下面请选择模板文件；选择后会自动记住。")
        self._choose_template_dialog()

    def _choose_template_dialog(self):
        """弹出文件框选择模板，选中返回 True。"""
        path = filedialog.askopenfilename(
            title="选择 Word 签字表模板",
            filetypes=[("Word 文件", "*.docx"), ("所有文件", "*.*")])
        if path:
            self.template_path.set(path)
            self._save_cfg()
            return True
        return False

    def _ensure_template(self):
        """模板缺失时弹窗引导；True 表示模板就绪。"""
        t = self.template_path.get()
        if t and os.path.exists(t):
            return True
        ans = messagebox.askyesnocancel(
            "需要 Word 模板",
            "未找到 Word 签字表模板。\n\n"
            "“是”立即选择模板文件；\n“否”稍后在界面中手动选择。")
        if ans is None:
            return False
        if ans:
            return self._choose_template_dialog()
        return False

    # ---------- 检查更新 ----------
    @staticmethod
    def _version_tuple(v):
        nums = re.findall(r"\d+", v or "")
        return tuple(int(x) for x in nums[:3]) + (0, 0, 0)

    def _check_update(self):
        """后台线程请求最新版本，失败静默，绝不打扰使用。"""
        def work():
            try:
                req = urllib.request.Request(
                    "https://api.github.com/repos/"
                    "yyxi746/moral-score-sign-generator/releases/latest",
                    headers={"Accept": "application/vnd.github+json",
                             "User-Agent": "moral-score-sign-generator"})
                with urllib.request.urlopen(req, timeout=6) as r:
                    data = json.loads(r.read().decode("utf-8"))
                tag = data.get("tag_name", "")
                url = data.get("html_url", "")
                if tag and self._version_tuple(tag) > self._version_tuple(APP_VERSION):
                    self.after(0, lambda: self._show_update(tag, url))
            except Exception:
                pass
        threading.Thread(target=work, daemon=True).start()

    def _show_update(self, tag, url):
        if getattr(self, "_update_shown", False):
            return
        if tag == getattr(self, "_skip_version", ""):
            return
        self._update_shown = True
        wrap = ttk.Frame(self._bottom)
        wrap.pack(side="right", padx=10)
        lbl = ttk.Label(wrap, text=f"发现新版本 {tag}，点击前往下载",
                        foreground="#1a6fe0", cursor="hand2")
        lbl.pack(side="left")
        lbl.bind("<Button-1>", lambda e: webbrowser.open(url))
        skip = ttk.Label(wrap, text="忽略此版本", foreground="gray",
                         cursor="hand2")
        skip.pack(side="left", padx=(8, 0))

        def do_skip(_=None):
            self._skip_version = tag
            self._save_cfg()
            wrap.destroy()

        skip.bind("<Button-1>", do_skip)

    def _build_ui(self):
        pad = {"padx": 8, "pady": 4}

        menubar = tk.Menu(self)
        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label="关于…", command=self.show_about)
        menubar.add_cascade(label="帮助", menu=help_menu, underline=0)
        self.config(menu=menubar)

        bottom = ttk.Frame(self)
        bottom.pack(side="bottom", fill="x", padx=10, pady=(0, 6))
        self._bottom = bottom
        self._make_info_button(bottom).pack(side="left")
        self.status = ttk.Label(bottom, text="就绪", foreground="gray")
        self.status.pack(side="left", padx=10)

        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True, padx=10, pady=(10, 4))
        self.nb.bind("<<NotebookTabChanged>>", self._on_tab_changed)

        root = ttk.Frame(self.nb)
        self.nb.add(root, text="单文件处理")
        batch = ttk.Frame(self.nb)
        self.nb.add(batch, text="批量处理")
        self.batch_tab = batch
        self._build_batch_tab(batch)

        r = 0
        ttk.Label(root, text="Excel 数据源：").grid(row=r, column=0, sticky="w", **pad)
        ttk.Entry(root, textvariable=self.excel_path).grid(
            row=r, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(root, text="浏览…", command=self.browse_excel).grid(row=r, column=4, **pad)

        r += 1
        ttk.Label(root, text="工作表：").grid(row=r, column=0, sticky="w", **pad)
        self.sheet_combo = ttk.Combobox(root, textvariable=self.sheet_name, state="readonly")
        self.sheet_combo.grid(row=r, column=1, columnspan=3, sticky="we", **pad)
        self.sheet_combo.bind("<<ComboboxSelected>>", lambda e: self.analyze())

        r += 1
        ttk.Label(root, text="Word 模板：").grid(row=r, column=0, sticky="w", **pad)
        ttk.Entry(root, textvariable=self.template_path).grid(
            row=r, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(root, text="浏览…", command=self.browse_template).grid(row=r, column=4, **pad)

        r += 1
        ttk.Label(root, text="Word 标题：").grid(row=r, column=0, sticky="w", **pad)
        ttk.Entry(root, textvariable=self.title_text).grid(
            row=r, column=1, columnspan=3, sticky="we", **pad)

        r += 1
        ttk.Label(root, text="班级：").grid(row=r, column=0, sticky="w", **pad)
        ttk.Entry(root, textvariable=self.class_name).grid(
            row=r, column=1, columnspan=3, sticky="we", **pad)

        r += 1
        ttk.Label(root, text="输出目录：").grid(row=r, column=0, sticky="w", **pad)
        ttk.Entry(root, textvariable=self.output_dir).grid(
            row=r, column=1, columnspan=2, sticky="we", **pad)
        ttk.Button(root, text="浏览…", command=self.browse_output_dir).grid(
            row=r, column=3, **pad)
        ttk.Button(root, text="清除", command=self.clear_output_dir).grid(
            row=r, column=4, **pad)

        r += 1
        ttk.Label(root, text="（输出目录留空 = 与数据源同目录）",
                  foreground="gray").grid(row=r, column=1, columnspan=3, sticky="w", padx=8)

        r += 1
        ttk.Label(root, text="输出文件：").grid(row=r, column=0, sticky="w", **pad)
        ttk.Entry(root, textvariable=self.output_path).grid(
            row=r, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(root, text="浏览…", command=self.browse_output).grid(row=r, column=4, **pad)

        r += 1
        map_frame = ttk.LabelFrame(root, text="列映射（已自动识别，可手动调整）")
        map_frame.grid(row=r, column=0, columnspan=5, sticky="we", padx=8, pady=6)
        ttk.Label(map_frame, text="姓名列：").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        self.name_combo = ttk.Combobox(map_frame, textvariable=self.name_label,
                                      state="readonly", width=24)
        self.name_combo.grid(row=0, column=1, sticky="w", padx=6)
        self.name_combo.bind("<<ComboboxSelected>>", lambda e: self.refresh_preview())

        ttk.Label(map_frame, text="德育分(总分)列：").grid(row=0, column=2, sticky="w", padx=6)
        self.score_combo = ttk.Combobox(map_frame, textvariable=self.score_label,
                                       state="readonly", width=24)
        self.score_combo.grid(row=0, column=3, sticky="w", padx=6)
        self.score_combo.bind("<<ComboboxSelected>>", lambda e: self.refresh_preview())

        ttk.Label(map_frame, text="表扬信\n(按月份段多选)：").grid(
            row=1, column=0, sticky="nw", padx=6, pady=6)
        self.praise_frame = ttk.Frame(map_frame)
        self.praise_frame.grid(row=1, column=1, columnspan=3, sticky="w", padx=6)

        r += 1
        prev_frame = ttk.LabelFrame(root, text="结果预览（生成前可核对，含全部学生）")
        prev_frame.grid(row=r, column=0, columnspan=5, sticky="nsew", padx=8, pady=6)
        self.tree = ttk.Treeview(prev_frame, columns=("name", "score", "praise"),
                                 show="headings", height=10)
        self.tree.heading("name", text="姓名")
        self.tree.heading("score", text="德育分")
        self.tree.heading("praise", text="表扬信")
        self.tree.column("name", width=220, anchor="center")
        self.tree.column("score", width=170, anchor="center")
        self.tree.column("praise", width=170, anchor="center")
        self.tree.tag_configure("warnrow", background="#fff3cd")
        self.tree.pack(side="left", fill="both", expand=True, padx=(6, 0), pady=6)
        sb = ttk.Scrollbar(prev_frame, orient="vertical", command=self.tree.yview)
        sb.pack(side="right", fill="y", pady=6)
        self.tree.configure(yscrollcommand=sb.set)

        r += 1
        self.count_label = ttk.Label(root, text="共识别到 0 人")
        self.count_label.grid(row=r, column=0, columnspan=5, sticky="w", padx=10)

        r += 1
        self.gen_btn = ttk.Button(root, text="生成签字表", command=self.generate)
        self.gen_btn.grid(row=r, column=0, columnspan=5, pady=10)

        root.columnconfigure(1, weight=1)
        root.rowconfigure(9, weight=1)

        # 文件拖放支持
        register_drop_recursive(root, self._on_drop_single)
        register_drop_recursive(batch, self._on_drop_batch)

    # ---------- 文件选择 ----------
    def browse_excel(self):
        try:
            path = filedialog.askopenfilename(
                title="选择德育分统计表",
                filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")])
            if path:
                self.excel_path.set(path)
                self.load_excel()
        except Exception as ex:
            self._error(ex)

    def browse_template(self):
        path = filedialog.askopenfilename(
            title="选择 Word 签字表模板",
            filetypes=[("Word 文件", "*.docx"), ("所有文件", "*.*")])
        if path:
            self.template_path.set(path)
            self._save_cfg()

    def browse_output_dir(self):
        d = filedialog.askdirectory(title="选择固定输出目录")
        if d:
            self.output_dir.set(d)
            self._save_cfg()
            self._recompute_output()

    def clear_output_dir(self):
        self.output_dir.set("")
        self._save_cfg()
        self._recompute_output()

    def browse_output(self):
        path = filedialog.asksaveasfilename(
            title="保存为", defaultextension=".docx",
            filetypes=[("Word 文件", "*.docx")])
        if path:
            self.output_path.set(path)

    # ---------- 加载 / 分析 ----------
    def load_excel(self):
        path = self.excel_path.get()
        if not os.path.exists(path):
            raise FileNotFoundError("所选 Excel 文件不存在。")
        if self.wb is not None:
            try:
                self.wb.close()
            except Exception:
                pass
        self.wb = openpyxl.load_workbook(path, data_only=True)
        names = list(self.wb.sheetnames)
        self.sheet_combo["values"] = names
        self.sheet_name.set(names[0] if names else "")
        self.analyze()

    def _current_ws(self):
        name = self.sheet_name.get()
        if not name or self.wb is None:
            return None
        return self.wb[name]

    def analyze(self):
        ws = self._current_ws()
        if ws is None:
            return
        self.headers = scan_headers(ws)
        self.col_options = column_labels(self.headers)
        self.label_to_col = {label: c for label, c in self.col_options}

        name_col, score_col, praise_cols = detect_columns(self.headers)
        labels_by_col = {c: label for label, c in self.col_options}
        self.name_label.set(labels_by_col.get(name_col, ""))
        self.score_label.set(labels_by_col.get(score_col, ""))

        self.month_groups = group_praise_by_month(ws, praise_cols)
        for w in self.praise_frame.winfo_children():
            w.destroy()
        self.month_vars = {}
        for i, (month, cols) in enumerate(self.month_groups):
            var = tk.IntVar(value=1)
            self.month_vars[month] = var
            ttk.Checkbutton(self.praise_frame, text=f"{month}（{len(cols)}列）",
                            variable=var,
                            command=self.refresh_preview).grid(
                row=i // 3, column=i % 3, sticky="w", padx=6, pady=3)

        self.title_text.set(extract_title(ws, self.sheet_name.get()))
        cls = extract_class(ws)
        if cls:
            self.class_name.set(cls)

        self._recompute_output()
        self.refresh_preview()

    def _recompute_output(self):
        excel = self.excel_path.get()
        if excel:
            self.output_path.set(compute_output_path(excel, self.output_dir.get()))

    def _selected_cols(self):
        name_col = self.label_to_col.get(self.name_label.get())
        score_col = self.label_to_col.get(self.score_label.get())
        praise_cols = []
        for month, cols in self.month_groups:
            if self.month_vars[month].get():
                praise_cols.extend(cols)
        return name_col, score_col, praise_cols

    def current_students(self):
        ws = self._current_ws()
        if ws is None:
            return []
        name_col, score_col, praise_cols = self._selected_cols()
        return extract_students(ws, name_col, score_col, praise_cols)

    def refresh_preview(self):
        try:
            students = self.current_students()
            for item in self.tree.get_children():
                self.tree.delete(item)
            problems = validate_students(students)
            bad = {p[1] for p in problems
                   if p[0] in ("empty", "neg", "badtext")}
            for i, st in enumerate(students):
                self.tree.insert(
                    "", "end",
                    values=(st["name"], st["score"], st["praise"]),
                    tags=("warnrow",) if i in bad else ())
            if problems and problems[0][0] == "zero":
                self.count_label.config(text="共识别到 0 人")
            elif problems:
                self.count_label.config(
                    text=f"共识别到 {len(students)} 人，"
                         f"{len(problems)} 项分数异常（黄色行）")
            else:
                self.count_label.config(text=f"共识别到 {len(students)} 人")
        except Exception as ex:
            self._error(ex)

    def _save_cfg(self):
        save_config({"template_path": self.template_path.get(),
                     "output_dir": self.output_dir.get().strip(),
                     "skip_version": getattr(self, "_skip_version", "")})

    # ---------- 批量处理页 ----------
    def _on_tab_changed(self, event=None):
        """批量页内容紧凑，自动降低窗口高度；切回单文件页恢复。"""
        try:
            if self.nb.select() == str(self.batch_tab):
                self.geometry("820x800")
            else:
                self.geometry("820x900")
        except Exception:
            pass

    def _build_batch_tab(self, parent):
        outer = ttk.Frame(parent)
        outer.pack(fill="both", expand=True, padx=8, pady=8)

        bar = ttk.Frame(outer)
        bar.pack(fill="x")
        ttk.Button(bar, text="添加文件",
                   command=self.batch_browse_files).pack(side="left", padx=4)
        ttk.Button(bar, text="添加文件夹",
                   command=self.batch_browse_folder).pack(side="left", padx=4)
        ttk.Button(bar, text="移除选中",
                   command=self.batch_remove).pack(side="left", padx=4)
        ttk.Button(bar, text="清空列表",
                   command=self.batch_clear).pack(side="left", padx=4)

        set_frame = ttk.Frame(outer)
        set_frame.pack(fill="x")
        ttk.Label(set_frame, text="Word 模板：").grid(
            row=0, column=0, sticky="w", padx=4, pady=3)
        ttk.Entry(set_frame, textvariable=self.template_path).grid(
            row=0, column=1, columnspan=3, sticky="we", padx=4)
        ttk.Button(set_frame, text="浏览…",
                   command=self.browse_template).grid(row=0, column=4, padx=4)
        ttk.Label(set_frame, text="输出目录：").grid(
            row=1, column=0, sticky="w", padx=4, pady=3)
        ttk.Entry(set_frame, textvariable=self.output_dir).grid(
            row=1, column=1, columnspan=2, sticky="we", padx=4)
        ttk.Button(set_frame, text="浏览…",
                   command=self.browse_output_dir).grid(row=1, column=3, padx=4)
        ttk.Button(set_frame, text="清除",
                   command=self.clear_output_dir).grid(row=1, column=4, padx=4)
        ttk.Label(set_frame, text="（输出目录留空 = 每个文件与数据源同目录）",
                  foreground="gray").grid(
            row=2, column=1, columnspan=3, sticky="w", padx=4)
        set_frame.columnconfigure(1, weight=1)

        month_frame = ttk.LabelFrame(
            outer, text="表扬信月份段（对所有文件生效，可多选）")
        month_frame.pack(fill="x")
        self.batch_praise_frame = ttk.Frame(month_frame)
        self.batch_praise_frame.pack(anchor="w", padx=8, pady=4)
        ttk.Label(self.batch_praise_frame,
                  text="添加文件并分析后，这里会显示可选月份段",
                  foreground="gray").grid(row=0, column=0, sticky="w")

        list_frame = ttk.LabelFrame(
            outer, text="文件列表（可直接把 Excel 文件或文件夹拖入窗口）")
        self.batch_tree = ttk.Treeview(
            list_frame, columns=("file", "sheet", "count", "status"),
            show="headings", height=8)
        self.batch_tree.heading("file", text="文件名")
        self.batch_tree.heading("sheet", text="工作表")
        self.batch_tree.heading("count", text="人数")
        self.batch_tree.heading("status", text="状态")
        self.batch_tree.column("file", width=280, anchor="w")
        self.batch_tree.column("sheet", width=140, anchor="center")
        self.batch_tree.column("count", width=60, anchor="center")
        self.batch_tree.column("status", width=260, anchor="w")
        self.batch_tree.tag_configure("error", foreground="#d23030")
        self.batch_tree.tag_configure("ok", foreground="#1a8f3c")
        self.batch_tree.tag_configure("warn", background="#fff3cd",
                                      foreground="#8a6d00")
        self.batch_tree.pack(side="left", fill="both", expand=True,
                             padx=(6, 0), pady=6)
        sb = ttk.Scrollbar(list_frame, orient="vertical",
                           command=self.batch_tree.yview)
        sb.pack(side="right", fill="y", pady=6)
        self.batch_tree.configure(yscrollcommand=sb.set)

        self.batch_tree.dnd_bind(
            "<<DragEnter>>",
            lambda e: self.status.config(text="松开鼠标即可添加文件…"))
        self.batch_tree.dnd_bind(
            "<<DragLeave>>", lambda e: self._batch_refresh_status())
        self.batch_tree.bind("<Double-1>", self._batch_edit_row)

        self.batch_gen_btn = ttk.Button(
            outer, text="开始批量生成", command=self.batch_generate)
        self.batch_gen_btn.pack(side="bottom", pady=4)
        self.batch_bar = ttk.Progressbar(outer, mode="determinate")
        self.batch_bar.pack(side="bottom", fill="x", pady=(6, 2))
        list_frame.pack(side="top", fill="both", expand=True, pady=8)

    def batch_browse_files(self):
        paths = filedialog.askopenfilenames(
            title="选择德育分统计表（可多选）",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")])
        if paths:
            self._batch_add_files(list(paths))

    def batch_browse_folder(self):
        d = filedialog.askdirectory(title="选择包含统计表的文件夹（自动递归）")
        if d:
            self._batch_add_files(collect_excels(d))

    def _batch_edit_row(self, event=None):
        item = self.batch_tree.identify_row(event.y) if event else None
        if not item or item not in self.batch_info:
            return
        info = self.batch_info[item]
        path = self.batch_items[item]
        ov = self.batch_overrides.get(item, {})
        top = tk.Toplevel(self)
        top.title("文件详情")
        top.transient(self)
        top.resizable(False, True)
        frm = ttk.Frame(top, padding=14)
        frm.pack(fill="both", expand=True)
        ttk.Label(frm, text=os.path.basename(path), foreground="gray").grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        ttk.Label(frm, text="Word 标题：").grid(row=1, column=0, sticky="w", pady=4)
        title_var = tk.StringVar(value=ov.get("title", info["title"]))
        ttk.Entry(frm, textvariable=title_var, width=48).grid(
            row=1, column=1, sticky="we", pady=4)
        ttk.Label(frm, text="班级：").grid(row=2, column=0, sticky="w", pady=4)
        class_var = tk.StringVar(value=ov.get("class", info["class"]))
        ttk.Entry(frm, textvariable=class_var, width=48).grid(
            row=2, column=1, sticky="we", pady=4)

        students = info["students"]
        detail = ttk.LabelFrame(
            frm, text=f"学生明细（{len(students)}人，黄色行=分数异常）")
        detail.grid(row=3, column=0, columnspan=2, sticky="nsew",
                    pady=(8, 0))
        dt = ttk.Treeview(detail, columns=("name", "score", "praise"),
                          show="headings", height=9)
        dt.heading("name", text="姓名")
        dt.heading("score", text="德育分")
        dt.heading("praise", text="表扬信(条)")
        dt.column("name", width=240, anchor="w")
        dt.column("score", width=100, anchor="center")
        dt.column("praise", width=100, anchor="center")
        dt.tag_configure("warnrow", background="#fff3cd")
        bad_rows = {p[1] for p in info.get("problems", [])
                    if p[0] in ("empty", "neg", "badtext")}
        for i, s in enumerate(students):
            dt.insert("", "end",
                      values=(s["name"], s["score"], s["praise"]),
                      tags=("warnrow",) if i in bad_rows else ())
        dt.pack(side="left", fill="both", expand=True, padx=(6, 0), pady=6)
        dsb = ttk.Scrollbar(detail, orient="vertical", command=dt.yview)
        dsb.pack(side="right", fill="y", pady=6)
        dt.configure(yscrollcommand=dsb.set)

        def status_text(suffix):
            probs = info.get("problems", [])
            if probs and probs[0][0] == "zero":
                return f"异常：没有学生数据（{suffix}）"
            if probs:
                return f"就绪（{len(probs)}项异常·{suffix}）"
            return f"就绪（{suffix}）"

        def save():
            t = title_var.get().strip()
            if not t:
                messagebox.showwarning("提示", "标题不能为空。", parent=top)
                return
            self.batch_overrides[item] = {"title": t,
                                         "class": class_var.get().strip()}
            cur = self.batch_tree.item(item, "values")
            tags = self.batch_tree.item(item, "tags")
            self.batch_tree.item(
                item,
                values=(cur[0], cur[1], cur[2], status_text("已修正")),
                tags=tags)
            top.destroy()

        def reset():
            self.batch_overrides.pop(item, None)
            cur = self.batch_tree.item(item, "values")
            tags = self.batch_tree.item(item, "tags")
            probs = info.get("problems", [])
            if probs and probs[0][0] == "zero":
                status = "异常：没有学生数据"
            elif probs:
                status = f"就绪（{len(probs)}项异常，双击查看）"
            else:
                status = "就绪"
            self.batch_tree.item(
                item, values=(cur[0], cur[1], cur[2], status), tags=tags)
            top.destroy()

        btns = ttk.Frame(frm)
        btns.grid(row=4, column=0, columnspan=2, pady=(10, 0))
        ttk.Button(btns, text="保存修正", command=save).pack(side="left", padx=6)
        ttk.Button(btns, text="恢复自动识别", command=reset).pack(side="left", padx=6)
        ttk.Button(btns, text="取消", command=top.destroy).pack(side="left", padx=6)
        top.grab_set()
        top.update_idletasks()
        top.geometry(f"+{self.winfo_rootx() + 60}+{self.winfo_rooty() + 30}")

    def batch_remove(self):
        for sel in self.batch_tree.selection():
            path = self.batch_items.pop(sel, None)
            self.batch_info.pop(sel, None)
            self.batch_overrides.pop(sel, None)
            if path:
                self.batch_path_set.discard(path)
            self.batch_tree.delete(sel)
        self._batch_refresh_status()

    def batch_clear(self):
        self.batch_items.clear()
        self.batch_info.clear()
        self.batch_path_set.clear()
        self.batch_overrides.clear()
        self._analyze_queue.clear()
        for i in self.batch_tree.get_children():
            self.batch_tree.delete(i)
        self._batch_refresh_status()

    def _batch_add_files(self, files):
        added = 0
        for f in files:
            f = os.path.normpath(f)
            if f in self.batch_path_set:
                continue
            self.batch_path_set.add(f)
            item = self.batch_tree.insert(
                "", "end",
                values=(os.path.basename(f), "-", "-", "等待分析…"))
            self.batch_items[item] = f
            self._analyze_queue.append(item)
            added += 1
        if added:
            self.status.config(text=f"已添加 {added} 个文件，正在分析…")
            self._pump_analyze()

    def _pump_analyze(self):
        if not self._analyze_queue:
            self._batch_refresh_status()
            return
        item = self._analyze_queue.pop(0)
        path = self.batch_items.get(item)
        try:
            info = analyze_excel_auto(path)
            problems = validate_students(info["students"])
            info["problems"] = problems
            self.batch_info[item] = info
            if problems and problems[0][0] == "zero":
                status, tags = "异常：没有学生数据", ("warn",)
            elif problems:
                status, tags = (f"就绪（{len(problems)}项异常，双击查看）",
                                ("warn",))
            else:
                status, tags = "就绪", ()
            self.batch_tree.item(
                item,
                values=(os.path.basename(path), info["sheet"],
                        len(info["students"]), status),
                tags=tags)
        except Exception as ex:
            self.batch_tree.item(
                item,
                values=(os.path.basename(path), "-", "-", f"失败：{ex}"),
                tags=("error",))
        if self._analyze_queue:
            self.after(10, self._pump_analyze)
        else:
            self._batch_refresh_status()

    def _refresh_batch_month_checks(self):
        """汇总所有已分析文件的月份段并集，生成复选框（保留旧勾选）。"""
        months, seen = [], set()
        for info in self.batch_info.values():
            for month, _cols in info.get("month_groups", []):
                if month not in seen:
                    seen.add(month)
                    months.append(month)
        months.sort(key=month_sort_key)
        old = self.batch_month_vars
        for w in self.batch_praise_frame.winfo_children():
            w.destroy()
        self.batch_month_vars = {}
        if not months:
            ttk.Label(self.batch_praise_frame,
                      text="添加文件并分析后，这里会显示可选月份段",
                      foreground="gray").grid(row=0, column=0, sticky="w")
            return
        for i, month in enumerate(months):
            was = old[month].get() if month in old else 1
            var = tk.IntVar(value=was)
            self.batch_month_vars[month] = var
            ttk.Checkbutton(self.batch_praise_frame, text=month,
                            variable=var).grid(
                row=i // 5, column=i % 5, sticky="w", padx=8, pady=3)

    def _batch_refresh_status(self):
        self._refresh_batch_month_checks()
        total = len(self.batch_items)
        ready = len(self.batch_info)
        if total == 0:
            self.status.config(text="就绪")
        else:
            self.status.config(text=f"文件 {total} 个，可处理 {ready} 个")

    def batch_generate(self):
        if not self._ensure_template():
            return
        ready_items = [it for it in self.batch_items if it in self.batch_info]
        if not ready_items:
            messagebox.showwarning("提示", "列表中没有可处理的文件。")
            return
        warn_items = [it for it in ready_items
                      if self.batch_info[it].get("problems")]
        if warn_items:
            ans = messagebox.askyesno(
                "数据异常提醒",
                f"有 {len(warn_items)} 个文件存在空分数、负分、非数字分数或无学生等问题。\n\n"
                "双击对应行可查看明细。\n是否仍要继续生成？")
            if not ans:
                return
        self._gen_queue = ready_items
        # 规划输出路径：不同子目录下的同名 Excel 自动加后缀，避免互相覆盖
        reserved = set()
        self._gen_path_map = {}
        for it in ready_items:
            p = compute_output_path(self.batch_items[it],
                                    self.output_dir.get())
            final = unique_output_path(p, reserved)
            reserved.add(final)
            self._gen_path_map[it] = final
        selected = {m for m, v in self.batch_month_vars.items() if v.get()}
        all_months = set(self.batch_month_vars)
        self._sel_months = None if selected == all_months else selected
        self._gen_results = {"ok": 0, "fail": 0, "errors": []}
        self._gen_cancel = False
        self.batch_bar["maximum"] = len(ready_items)
        self.batch_bar["value"] = 0
        self.batch_gen_btn.config(text="取消生成",
                                  command=self._cancel_generate,
                                  state="normal")
        self.status.config(text="开始批量生成…")
        self._pump_generate()

    def _cancel_generate(self):
        self._gen_cancel = True
        self.status.config(text="正在取消，完成当前文件后停止…")

    def _pump_generate(self):
        if getattr(self, "_gen_cancel", False):
            self._batch_finish(cancelled=True)
            return
        if not self._gen_queue:
            self._batch_finish()
            return
        item = self._gen_queue.pop(0)
        path = self.batch_items[item]
        cur = self.batch_tree.item(item, "values")
        try:
            out, n = process_one_excel(
                path, self.template_path.get(), self.output_dir.get(),
                self._sel_months, self.batch_overrides.get(item),
                self._gen_path_map.get(item))
            self._gen_results["ok"] += 1
            self.batch_tree.item(
                item,
                values=(cur[0], cur[1], cur[2], f"成功（{n}人）：{out}"),
                tags=("ok",))
        except Exception as ex:
            self._gen_results["fail"] += 1
            self._gen_results["errors"].append(
                (os.path.basename(path), str(ex)))
            self.batch_tree.item(
                item,
                values=(cur[0], cur[1], cur[2], f"失败：{ex}"),
                tags=("error",))
        self.batch_bar["value"] += 1
        self.status.config(
            text=f"已处理 {int(self.batch_bar['value'])}/{self.batch_bar['maximum']}")
        self.after(10, self._pump_generate)

    def _batch_finish(self, cancelled=False):
        self.batch_gen_btn.config(text="开始批量生成",
                                  command=self.batch_generate,
                                  state="normal")
        r = self._gen_results
        if cancelled:
            done = r["ok"] + r["fail"]
            total = int(self.batch_bar["maximum"])
            self.status.config(
                text=f"已取消：完成 {done}/{total}，成功 {r['ok']} 个，"
                     f"失败 {r['fail']} 个")
        else:
            self.status.config(
                text=f"批量完成：成功 {r['ok']} 个，失败 {r['fail']} 个")
        top = tk.Toplevel(self)
        top.title("批量生成已取消" if cancelled else "批量生成完成")
        top.transient(self)
        top.grab_set()
        top.geometry(
            f"440x280+{self.winfo_x()+200}+{self.winfo_y()+240}")
        head = ("批量生成已取消，" if cancelled else "")
        ttk.Label(top,
                  text=f"{head}成功 {r['ok']} 个，失败 {r['fail']} 个",
                  font=("Microsoft YaHei UI", 13, "bold")).pack(
            pady=(20, 8))
        if r["errors"]:
            msg = "\n".join(f"{n}：{e}" for n, e in r["errors"][:6])
            ttk.Label(top, text=msg, foreground="#d23030",
                      justify="left").pack(padx=20, pady=4)
        btns = ttk.Frame(top)
        btns.pack(pady=14)
        out_dir = self.output_dir.get().strip()
        if out_dir:
            ttk.Button(
                btns, text="打开输出目录",
                command=lambda: (top.destroy(), os.startfile(out_dir))
            ).grid(row=0, column=0, padx=8)
        ttk.Button(btns, text="关闭", command=top.destroy).grid(
            row=0, column=1, padx=8)

    # ---------- 拖放 ----------
    def _parse_drop_paths(self, data):
        try:
            items = list(self.tk.splitlist(data))
        except Exception:
            items = [data]
        files = []
        for it in items:
            it = str(it).strip().strip("{}")
            if not it:
                continue
            if os.path.isdir(it):
                files.extend(collect_excels(it))
            elif it.lower().endswith((".xlsx", ".xlsm")):
                files.append(it)
        return files

    def _on_drop_single(self, event):
        files = self._parse_drop_paths(event.data)
        if not files:
            return
        if len(files) == 1:
            self.excel_path.set(files[0])
            self.load_excel()
        else:
            self.nb.select(self.batch_tab)
            self._batch_add_files(files)

    def _on_drop_batch(self, event):
        files = self._parse_drop_paths(event.data)
        self._batch_add_files(files)

    # ---------- 生成 ----------
    def generate(self):
        try:
            excel = self.excel_path.get()
            template = self.template_path.get()
            output = self.output_path.get()
            if not excel or not os.path.exists(excel):
                messagebox.showwarning("提示", "请先选择有效的 Excel 文件。")
                return
            if not self._ensure_template():
                return
            name_col, score_col, _ = self._selected_cols()
            if not name_col:
                messagebox.showwarning("提示", "请指定“姓名列”。")
                return
            if not score_col:
                messagebox.showwarning("提示", "请指定“德育分(总分)列”。")
                return
            students = self.current_students()
            if not students:
                messagebox.showwarning("提示", "没有识别到任何学生数据，请检查列映射。")
                return
            problems = validate_students(students)
            if problems:
                parts = []
                for p in problems[:5]:
                    if p[0] == "empty":
                        parts.append(f"{p[2]} 分数为空")
                    elif p[0] == "neg":
                        parts.append(f"{p[2]} 分数为负({p[3]})")
                    else:
                        parts.append(f"{p[2]} 分数非数字({p[3]})")
                more = "…" if len(problems) > 5 else ""
                ans = messagebox.askyesno(
                    "数据异常提醒",
                    f"检测到 {len(problems)} 项异常：\n"
                    + "；".join(parts) + more + "\n\n是否仍要生成？")
                if not ans:
                    return
            out_dir = os.path.dirname(output)
            if not output or not out_dir:
                messagebox.showwarning("提示", "请指定输出文件路径。")
                return
            if not os.path.isdir(out_dir):
                messagebox.showwarning("提示", f"输出目录不存在：\n{out_dir}")
                return
            if os.path.exists(output):
                if not messagebox.askyesno(
                        "文件已存在",
                        f"输出文件已存在：\n{os.path.basename(output)}\n\n"
                        "是否覆盖？"):
                    return

            self.gen_btn.config(state="disabled")
            self.status.config(text="正在生成…")
            self.update_idletasks()
            fill_word(template, output, students,
                      self.title_text.get(), self.class_name.get())
            self.status.config(text=f"已生成：{output}")
            self._ask_open(output, len(students))
        except Exception as ex:
            self._error(ex)
        finally:
            try:
                self.gen_btn.config(state="normal")
            except Exception:
                pass

    # ---------- 关于 ----------
    def _make_info_button(self, parent):
        """左下角圆形 ⓘ 按钮，hover 变蓝，点击弹出"关于"窗口。"""
        bg = ttk.Style().lookup("TFrame", "background") or "#f0f0f0"
        size = 34
        cx = size / 2
        cv = tk.Canvas(parent, width=size, height=size, bg=bg,
                       highlightthickness=0, bd=0, cursor="hand2")
        ring = cv.create_oval(3, 3, size - 3, size - 3,
                              outline="#9aa4b2", width=1.5)
        dot = cv.create_oval(cx - 1.8, 9.5, cx + 1.8, 13.1,
                             fill="#9aa4b2", outline="")
        bar = cv.create_line(cx, 15.5, cx, 23.5, fill="#9aa4b2",
                             width=1.8, capstyle="round")

        def set_color(c):
            cv.itemconfigure(ring, outline=c)
            cv.itemconfigure(dot, fill=c)
            cv.itemconfigure(bar, fill=c)

        cv.bind("<Enter>", lambda e: set_color("#2b7fff"))
        cv.bind("<Leave>", lambda e: set_color("#9aa4b2"))
        cv.bind("<Button-1>", lambda e: self.show_about())
        return cv

    def show_about(self):
        top = tk.Toplevel(self)
        top.title("关于")
        top.transient(self)
        top.grab_set()
        top.resizable(False, False)
        top.geometry(f"450x330+{self.winfo_x()+180}+{self.winfo_y()+240}")

        ttk.Label(top, text="德育分签字表生成器",
                  font=("Microsoft YaHei UI", 15, "bold")).pack(pady=(22, 4))
        ttk.Label(top, text=f"版本 {APP_VERSION}").pack()
        ttk.Label(top,
                  text="根据 Excel 德育分统计表，自动生成 Word 德育分确认签字表。").pack(pady=(10, 2))
        ttk.Label(top, text="Python · tkinter · openpyxl · python-docx",
                  foreground="gray").pack()
        link = ttk.Label(top, text=GITHUB_URL, foreground="#0066cc", cursor="hand2")
        link.pack(pady=12)
        link.bind("<Button-1>", lambda e: webbrowser.open(GITHUB_URL))

        btns = ttk.Frame(top)
        btns.pack(pady=6)

        def copy_link():
            self.clipboard_clear()
            self.clipboard_append(GITHUB_URL)
            self.status.config(text="GitHub 链接已复制到剪贴板")

        ttk.Button(btns, text="访问 GitHub",
                   command=lambda: webbrowser.open(GITHUB_URL)).grid(row=0, column=0, padx=6)
        ttk.Button(btns, text="复制链接", command=copy_link).grid(row=0, column=1, padx=6)
        ttk.Button(btns, text="关闭", command=top.destroy).grid(row=0, column=2, padx=6)

    def _safe_start(self, target, top=None):
        try:
            os.startfile(target)
            if top is not None:
                top.destroy()
        except Exception as ex:
            self._error(ex)

    def _ask_open(self, output, n):
        top = tk.Toplevel(self)
        top.title("生成成功")
        top.transient(self)
        top.grab_set()
        top.geometry(f"360x150+{self.winfo_x()+200}+{self.winfo_y()+220}")
        ttk.Label(top, text=f"已生成，共 {n} 人。").pack(pady=(24, 8))
        btns = ttk.Frame(top)
        btns.pack(pady=8)
        ttk.Button(btns, text="打开文件",
                   command=lambda: self._safe_start(output, top)).grid(
            row=0, column=0, padx=8)
        ttk.Button(btns, text="打开文件夹",
                   command=lambda: self._safe_start(
                       os.path.dirname(output), top)).grid(
            row=0, column=1, padx=8)
        ttk.Button(btns, text="关闭", command=top.destroy).grid(
            row=0, column=2, padx=8)

    def _error(self, ex):
        try:
            self.status.config(text="出错了")
            messagebox.showerror("出错了", f"{ex}")
        except Exception:
            pass


def main():
    App().mainloop()


if __name__ == "__main__":
    main()
